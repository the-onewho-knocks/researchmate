"""backend/app/services/extraction/docx_extractor.py"""
import re
from pathlib import Path

from docx import Document

from app.core.errors import CorruptFileError, ExtractionError
from app.schemas.document import Block, ParsedDocument

from .base import Extractor
from .normalizer import clean_text


class DocxExtractor(Extractor):
    def extract(self, path: Path) -> ParsedDocument:
        try:
            doc = Document(str(path))
        except Exception as exc:
            raise CorruptFileError(f"Could not open the .docx file: {exc}") from exc

        blocks: list[Block] = []
        for p in doc.paragraphs:
            text = clean_text(p.text)
            if not text:
                continue

            style = ((p.style.name if p.style is not None else "") or "").strip().lower()
            kind, level = "text", None
            if style == "title":
                kind = "title"
            elif style.startswith("heading"):
                kind = "heading"
                digits = re.sub(r"\D", "", style)
                level = int(digits) if digits else 1
            elif style == "caption":
                kind = "caption"

            runs = [r for r in p.runs if r.text.strip()]
            size, bold = None, False
            if runs:
                if runs[0].font.size:
                    size = float(runs[0].font.size.pt)
                bold = all(bool(r.bold) for r in runs)

            blocks.append(Block(text=text, size=size, bold=bold, kind=kind, level=level))

        if not blocks:
            raise ExtractionError("The .docx file contains no text.")

        return ParsedDocument(
            file_type="docx",
            blocks=blocks,
            num_pages=None,   # Word does not store a reliable page count
            meta={"n_tables": len(doc.tables), "n_figures": len(doc.inline_shapes), "n_images": len(doc.inline_shapes)},
        )
